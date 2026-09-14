"""
core/index_integrity.py — Processing-state audit for ticket #3.

First cut (Cursor / CI):
  Account for every on-disk file in a *synthetic* archive as one of:
    indexed | metadata_only | excluded | duplicate | failed
  INDEXED status alone is not treated as content-quality proof.

  Emits machine-readable JSON + a readable summary with build SHA.
  Does **not** claim office-corpus acceptance (that stays on WOHA/Codex).
  Does **not** score cross-project contamination, retrieval top-3, or synthesis.
"""

from __future__ import annotations

import json
import logging
import os
import sqlite3
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from unittest.mock import patch

import yaml

from config import Config, load_config
from core.db import get_connection
from core.discover import _compute_fingerprint_by_strategy as _orig_fingerprint
from core.discover import _is_excluded
from core.index import run_full_pipeline
from core.pipeline_validate import mock_embed_if_needed

logger = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXPECTED = REPO_ROOT / "tests" / "fixtures" / "integrity_expected.yaml"

PROCESSING_STATES = (
    "indexed",
    "metadata_only",
    "excluded",
    "duplicate",
    "failed",
)

# Filename the synthetic fixture uses to exercise discover READ_ERROR.
FORCE_READ_ERROR_NAME = "unreadable.bin"

_OFFICE_NOTE = (
    "Real-corpus acceptance (NUS / NParks / HICA on WOHA) stays with local Codex. "
    "This run is synthetic CI only and must not be treated as an office gate."
)


def get_build_sha(repo_root: Path | None = None) -> str:
    """Return the current git HEAD SHA, or 'unknown' if git is unavailable."""
    root = repo_root or REPO_ROOT
    try:
        out = subprocess.check_output(
            ["git", "rev-parse", "HEAD"],
            cwd=root,
            stderr=subprocess.DEVNULL,
            text=True,
        ).strip()
        return out or "unknown"
    except Exception:
        return "unknown"


def _iter_all_files(root: Path):
    """Walk every regular file, including paths that exclude_globs would drop."""
    stack = [root]
    while stack:
        current = stack.pop()
        try:
            with os.scandir(current) as it:
                for entry in it:
                    try:
                        if entry.is_dir(follow_symlinks=False):
                            stack.append(Path(entry.path))
                        elif entry.is_file(follow_symlinks=False):
                            yield Path(entry.path)
                    except OSError:
                        continue
        except (PermissionError, OSError):
            continue


def _rel_posix(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _chunk_has_real_text(ref_value: str, text: str) -> bool:
    """True when a chunk is extracted content, not a path-only surrogate."""
    if not (text or "").strip():
        return False
    if (ref_value or "") == "meta":
        return False
    return True


def classify_file(
    *,
    rel_path: str,
    on_disk: bool,
    excluded_by_glob: bool,
    row: dict[str, Any] | None,
    chunks: list[dict[str, Any]],
) -> dict[str, Any]:
    """
    Map disk + DB facts to one explainable processing state.

    INDEXED / EMBEDDED / EXTRACTED is not enough: text-extractable files
    without real extracted text are metadata_only, with missing_text visible.
    """
    status = (row or {}).get("status")
    lane = (row or {}).get("lane")
    error_code = (row or {}).get("error_code")
    error_detail = (row or {}).get("error_detail")
    duplicate_of = (row or {}).get("duplicate_of")

    real_chunks = [
        c for c in chunks
        if _chunk_has_real_text(c.get("ref_value", ""), c.get("text", ""))
    ]
    missing_text = (bool(chunks) and not real_chunks) or (
        not chunks and status in {"INDEXED", "EMBEDDED", "EXTRACTED", "FAILED"}
    )
    stale_embeddings = bool(
        status == "INDEXED"
        and chunks
        and any(not c.get("embedded") for c in chunks)
    )

    record: dict[str, Any] = {
        "rel_path": rel_path,
        "on_disk": on_disk,
        "db_status": status,
        "lane": lane,
        "error_code": error_code,
        "error_detail": error_detail,
        "duplicate_of": duplicate_of,
        "chunk_count": len(chunks),
        "real_text_chunks": len(real_chunks),
        "missing_text": missing_text,
        "stale_embeddings": stale_embeddings,
        "failure_class": None,
        "reason": "",
        "processing_state": "unexplained",
    }

    if excluded_by_glob:
        record.update(
            processing_state="excluded",
            reason="exclude_glob: path matches configured exclude_globs",
        )
        return record

    if row is None:
        record.update(
            processing_state="unexplained",
            reason="on disk, not excluded, and not present in the files table",
        )
        return record

    if status == "FAILED":
        record.update(
            processing_state="failed",
            failure_class="extraction",
            reason=f"extraction failure: {error_code or 'FAILED'}"
            + (f" ({error_detail})" if error_detail else ""),
        )
        return record

    if status == "SKIPPED" and error_code == "DUPLICATE":
        record.update(
            processing_state="duplicate",
            reason=f"content duplicate of {duplicate_of or error_detail or 'canonical file'}",
        )
        return record

    if status == "SKIPPED":
        record.update(
            processing_state="excluded",
            reason=f"policy skip: {error_code or 'SKIPPED'}",
        )
        return record

    if lane == "METADATA_ONLY":
        record.update(
            processing_state="metadata_only",
            reason="intentional metadata-only lane (CAD / BIM / media / unknown ext)",
            missing_text=True,
        )
        return record

    if error_code == "EXTRACT_EMPTY_FALLBACK":
        record.update(
            processing_state="metadata_only",
            reason="no extractable text — path/filename surrogate only (EXTRACT_EMPTY_FALLBACK)",
            missing_text=True,
        )
        return record

    if status in {"INDEXED", "EMBEDDED", "EXTRACTED"} and real_chunks:
        record.update(
            processing_state="indexed",
            reason=f"extracted text present ({len(real_chunks)} content chunk(s)); status={status}",
            missing_text=False,
        )
        return record

    if status in {"INDEXED", "EMBEDDED", "EXTRACTED"}:
        record.update(
            processing_state="metadata_only",
            reason="pipeline status is INDEXED/EXTRACTED but no real extracted text — INDEXED is not content proof",
            missing_text=True,
        )
        return record

    if status == "DISCOVERED":
        record.update(
            processing_state="unexplained",
            reason="still DISCOVERED — extract/index did not complete",
        )
        return record

    record.update(
        processing_state="unexplained",
        reason=f"no mapping for status={status!r} lane={lane!r} error_code={error_code!r}",
    )
    return record


def _load_db_files(conn: sqlite3.Connection) -> dict[str, dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT file_id, file_path, file_name, extension, lane, status,
               error_code, error_detail, duplicate_of
        FROM files
        """
    ).fetchall()
    return {r["file_path"]: dict(r) for r in rows}


def _load_chunks_by_file(conn: sqlite3.Connection) -> dict[str, list[dict[str, Any]]]:
    rows = conn.execute(
        "SELECT file_id, ref_value, text, embedded FROM chunks"
    ).fetchall()
    out: dict[str, list[dict[str, Any]]] = {}
    for r in rows:
        out.setdefault(r["file_id"], []).append(dict(r))
    return out


def audit_processing_states(
    conn: sqlite3.Connection,
    archive_root: Path,
    cfg: Config,
) -> list[dict[str, Any]]:
    """Classify every on-disk file and any leftover DB-only rows."""
    archive_root = archive_root.resolve()
    by_path = _load_db_files(conn)
    chunks_by_file = _load_chunks_by_file(conn)
    seen_db_paths: set[str] = set()
    records: list[dict[str, Any]] = []

    for path in _iter_all_files(archive_root):
        posix = path.resolve().as_posix()
        row = by_path.get(posix)
        if row:
            seen_db_paths.add(posix)
        chunks = chunks_by_file.get((row or {}).get("file_id", ""), []) if row else []
        records.append(
            classify_file(
                rel_path=_rel_posix(path, archive_root),
                on_disk=True,
                excluded_by_glob=_is_excluded(path, cfg.exclude_globs),
                row=row,
                chunks=chunks,
            )
        )

    for posix, row in by_path.items():
        if posix in seen_db_paths:
            continue
        chunks = chunks_by_file.get(row.get("file_id", ""), [])
        records.append(
            classify_file(
                rel_path=_rel_posix(Path(posix), archive_root),
                on_disk=False,
                excluded_by_glob=False,
                row=row,
                chunks=chunks,
            )
        )

    records.sort(key=lambda r: r["rel_path"])
    return records


def build_synthetic_archive(dest: Path) -> Path:
    """
    Write a small synthetic architectural tree covering the five processing states.

    Names are NUS-like only as a stand-in. No office / WOHA source files.
    """
    dest = dest.resolve()
    if dest.exists():
        import shutil
        shutil.rmtree(dest)
    dest.mkdir(parents=True)

    root = dest / "101_nus_synth"
    brief = (
        "# 101 NUS Synth — Design Brief\n\n"
        "Synthetic campus briefing note for index-integrity CI.\n"
        "Client asked for 1,200 undergraduate places and a shared lecture theatre.\n"
        "This is not a real NUS or WOHA project file.\n"
    )
    (root / "01_Reports").mkdir(parents=True)
    (root / "01_Reports" / "design_brief.txt").write_text(brief, encoding="utf-8")

    (root / "copies").mkdir(parents=True)
    (root / "copies" / "design_brief.txt").write_text(brief, encoding="utf-8")

    (root / "04_Spreadsheets").mkdir(parents=True)
    (root / "04_Spreadsheets" / "area_schedule.csv").write_text(
        "space,gfa_sqm,notes\n"
        "lecture_theatre,850,shared teaching\n"
        "studio,1200,architecture studio\n",
        encoding="utf-8",
    )

    (root / "05_Correspondence").mkdir(parents=True)
    (root / "05_Correspondence" / "client_reply.eml").write_text(
        "From: pm@example.com\n"
        "To: client@nus-synth.example\n"
        "Subject: 101 NUS Synth — GFA clarification\n"
        "Date: Mon, 1 Jan 2024 10:00:00 +0800\n"
        "\n"
        "Please confirm the lecture theatre GFA remains 850 sqm.\n",
        encoding="utf-8",
    )

    (root / "06_Drawings").mkdir(parents=True)
    (root / "06_Drawings" / "level_01.dwg").write_bytes(b"AC10 synthetic dwg placeholder\n")

    (root / "07_Models").mkdir(parents=True)
    (root / "07_Models" / "massing.rvt").write_bytes(b"RVT synthetic model placeholder\n")

    (root / "junk").mkdir(parents=True)
    (root / "junk" / "~$lock.docx").write_bytes(b"office lock file")
    (root / "junk" / "scratch.tmp").write_text("temp scratch", encoding="utf-8")

    (root / "failed").mkdir(parents=True)
    (root / "failed" / FORCE_READ_ERROR_NAME).write_bytes(b"should not be readable during discover")

    (root / "02_Presentations").mkdir(parents=True)
    _write_pptx(
        root / "02_Presentations" / "board_update.pptx",
        title="101 NUS Synth board update",
        body="Campus massing option B and lecture theatre GFA.",
    )

    (root / "03_Scanned").mkdir(parents=True)
    _write_blank_pdf(root / "03_Scanned" / "site_scan.pdf")

    return dest


def _write_pptx(path: Path, title: str, body: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(4))
    tf = box.text_frame
    tf.text = title
    p = tf.add_paragraph()
    p.text = body
    prs.save(str(path))


def _write_blank_pdf(path: Path) -> None:
    import pypdf

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as fh:
        writer.write(fh)


def _fingerprint_with_forced_read_errors(path_str: str, strategy: str, size_bytes: int, mtime_epoch: float):
    if Path(path_str).name == FORCE_READ_ERROR_NAME:
        raise OSError("synthetic fixture: forced READ_ERROR")
    return _orig_fingerprint(path_str, strategy, size_bytes, mtime_epoch)


def setup_integrity_sandbox(work_dir: Path) -> tuple[Config, Path]:
    """Isolated tiga_work + synthetic archive. Full fingerprints so dupes are visible."""
    work_dir = work_dir.resolve()
    work_dir.mkdir(parents=True, exist_ok=True)
    archive_root = work_dir / "integrity_archive"
    build_synthetic_archive(archive_root)

    cfg_file = work_dir / "config.yaml"
    cfg_data = {
        "index_roots": [str(archive_root)],
        "retrieval": {"top_k_default": 10, "reranker_enabled": False},
        "compose": {"api_enabled": False, "provider": "ollama"},
        "pipeline": {
            "extract_workers": 1,
            "fingerprint_strategy": "full",
        },
        "dedupe": {"enabled": True},
        "extract": {"empty_fallback_metadata": True},
        "scheduler": {
            "day_run_indexing": True,
            "night_run_indexing": True,
        },
    }
    cfg_file.write_text(yaml.dump(cfg_data), encoding="utf-8")
    cfg = load_config(config_file=cfg_file, work_dir=work_dir)
    cfg.ensure_dirs()
    return cfg, archive_root


def _counts(records: list[dict[str, Any]]) -> dict[str, int]:
    counts = {state: 0 for state in PROCESSING_STATES}
    counts["unexplained"] = 0
    for rec in records:
        state = rec.get("processing_state", "unexplained")
        counts[state] = counts.get(state, 0) + 1
    return counts


def _match_expected(rel_path: str, expected_suffix: str) -> bool:
    rel = rel_path.replace("\\", "/")
    exp = expected_suffix.replace("\\", "/").lstrip("./")
    return rel == exp or rel.endswith("/" + exp) or rel.endswith(exp)


def compare_to_expected(
    records: list[dict[str, Any]],
    expected_path: Path | None = None,
) -> dict[str, Any]:
    """Score classified files against the versioned expected-state manifest."""
    path = expected_path or DEFAULT_EXPECTED
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) if path.exists() else {}
    expected_files = list(raw.get("files") or [])

    mismatches: list[dict[str, Any]] = []
    matched = 0
    pairs: dict[str, list[dict[str, Any]]] = {}
    singles: list[dict[str, Any]] = []
    for entry in expected_files:
        if entry.get("pair"):
            pairs.setdefault(str(entry["pair"]), []).append(entry)
        else:
            singles.append(entry)

    def _lookup(suffix: str) -> dict[str, Any] | None:
        found = [r for r in records if _match_expected(r["rel_path"], suffix)]
        return found[0] if found else None

    for entry in singles:
        suffix = entry.get("path", "")
        want = entry.get("processing_state")
        rec = _lookup(suffix)
        if rec is None:
            mismatches.append({"path": suffix, "expected": want, "actual": None, "error": "missing"})
            continue
        if rec["processing_state"] != want:
            mismatches.append({
                "path": suffix,
                "expected": want,
                "actual": rec["processing_state"],
                "error": "state_mismatch",
                "reason": rec.get("reason"),
            })
        else:
            matched += 1

    # Duplicate pairs: walk order can flip which copy is canonical.
    for pair_id, entries in pairs.items():
        actual_states: list[str] = []
        missing_paths: list[str] = []
        for entry in entries:
            rec = _lookup(entry.get("path", ""))
            if rec is None:
                missing_paths.append(entry.get("path", ""))
            else:
                actual_states.append(rec["processing_state"])
        want_states = sorted(e.get("processing_state") for e in entries)
        if missing_paths:
            mismatches.append({
                "path": pair_id,
                "expected": want_states,
                "actual": actual_states,
                "error": "missing",
                "missing_paths": missing_paths,
            })
            continue
        if sorted(actual_states) != want_states:
            mismatches.append({
                "path": pair_id,
                "expected": want_states,
                "actual": actual_states,
                "error": "pair_state_mismatch",
            })
        else:
            matched += len(entries)

    return {
        "expected_path": str(path),
        "expected_count": len(expected_files),
        "matched": matched,
        "mismatches": mismatches,
        "ok": len(mismatches) == 0 and len(expected_files) > 0,
    }


def format_summary(report: dict[str, Any]) -> str:
    """Human-readable summary. Never claims office-corpus acceptance."""
    cov = report.get("coverage") or {}
    by_state = cov.get("by_state") or {}
    lat = report.get("latency_ms") or {}
    lines = [
        "TIGA index integrity — synthetic fixture audit",
        f"Ticket          : #{report.get('ticket', '3')}",
        f"Build SHA       : {report.get('build_sha', 'unknown')}",
        f"Corpus          : {(report.get('corpus_snapshot') or {}).get('name', 'synthetic')} "
        f"({cov.get('files_on_disk', 0)} files on disk)",
        f"Office corpus   : NOT CLAIMED — {report.get('office_corpus_note', _OFFICE_NOTE)}",
        f"Contamination   : not evaluated (out of scope for this cut)",
        f"Retrieval/synth : not evaluated (extraction-state audit only)",
        "",
        f"Coverage        : {cov.get('accounted', 0)}/{cov.get('files_on_disk', 0)} accounted"
        f"  unexplained={cov.get('unexplained', 0)}",
        f"  indexed         : {by_state.get('indexed', 0)}",
        f"  metadata_only   : {by_state.get('metadata_only', 0)}",
        f"  excluded        : {by_state.get('excluded', 0)}",
        f"  duplicate       : {by_state.get('duplicate', 0)}",
        f"  failed          : {by_state.get('failed', 0)}",
        "",
        f"Latency         : index {lat.get('index', 0):.0f} ms | "
        f"audit {lat.get('audit', 0):.0f} ms | total {lat.get('total', 0):.0f} ms",
        f"Expected states : {'PASS' if (report.get('expected') or {}).get('ok') else 'FAIL'}"
        f"  ({(report.get('expected') or {}).get('matched', 0)}/"
        f"{(report.get('expected') or {}).get('expected_count', 0)} matched)",
        f"Audit           : {'PASS' if report.get('ok') else 'FAIL'}",
    ]
    for rec in report.get("files") or []:
        if rec.get("processing_state") in {"failed", "unexplained"} or rec.get("stale_embeddings"):
            lines.append(
                f"  - {rec.get('processing_state')}: {rec.get('rel_path')} "
                f"({rec.get('reason')})"
            )
    return "\n".join(lines) + "\n"


def write_reports(report: dict[str, Any], report_dir: Path) -> tuple[Path, Path]:
    report_dir.mkdir(parents=True, exist_ok=True)
    ts = report.get("ts") or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
    json_path = report_dir / f"integrity_{ts}.json"
    txt_path = report_dir / f"integrity_{ts}.txt"
    json_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    txt_path.write_text(format_summary(report), encoding="utf-8")
    return json_path, txt_path


def run_integrity_benchmark(
    work_dir: Path | None = None,
    *,
    expected_path: Path | None = None,
    mock_embed: bool = True,
    verbose: bool = True,
) -> dict[str, Any]:
    """
    Build synthetic archive → index → processing-state audit → write reports.

    Office-corpus acceptance is always false.
    """
    sandbox = (work_dir or (REPO_ROOT / "tiga_work_integrity")).resolve()
    cfg, archive_root = setup_integrity_sandbox(sandbox)
    conn = get_connection(cfg.get_db_path())
    t0 = time.perf_counter()
    try:
        with mock_embed_if_needed(mock_embed), patch(
            "core.discover._compute_fingerprint_by_strategy",
            side_effect=_fingerprint_with_forced_read_errors,
        ):
            t_index = time.perf_counter()
            index_stats = run_full_pipeline(conn, cfg)
            index_ms = round((time.perf_counter() - t_index) * 1000, 1)

            t_audit = time.perf_counter()
            records = audit_processing_states(conn, archive_root, cfg)
            audit_ms = round((time.perf_counter() - t_audit) * 1000, 1)

        expected = compare_to_expected(records, expected_path)
        counts = _counts(records)
        unexplained = [r for r in records if r["processing_state"] == "unexplained"]
        on_disk = sum(1 for r in records if r.get("on_disk"))
        accounted = sum(1 for r in records if r["processing_state"] in PROCESSING_STATES)
        errors: list[str] = []
        if unexplained:
            errors.append(f"{len(unexplained)} file(s) have no explainable processing state")
        if not expected["ok"]:
            errors.append(f"{len(expected['mismatches'])} expected-state mismatch(es)")

        sha = get_build_sha()
        ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        total_ms = round((time.perf_counter() - t0) * 1000, 1)
        report = {
            "ticket": "3",
            "title": "Index integrity (synthetic fixtures)",
            "ts": ts,
            "office_corpus_acceptance": False,
            "office_corpus_note": _OFFICE_NOTE,
            "scope": {
                "corpus": "synthetic_fixtures",
                "cross_project_contamination": "not_in_scope",
                "retrieval_top3_gate": "not_claimed",
                "synthesis_factuality_gate": "not_claimed",
                "failure_classes_evaluated": ["extraction"],
            },
            "build_sha": sha,
            "build_sha_short": sha[:12] if sha != "unknown" else "unknown",
            "corpus_snapshot": {
                "kind": "synthetic",
                "name": "integrity_archive",
                "root": str(archive_root),
                "file_count_on_disk": on_disk,
                "fingerprint_strategy": "full",
            },
            "latency_ms": {
                "index": index_ms,
                "audit": audit_ms,
                "total": total_ms,
            },
            "coverage": {
                "files_on_disk": on_disk,
                "accounted": accounted,
                "unexplained": len(unexplained),
                "by_state": counts,
            },
            "index_stats": index_stats,
            "expected": expected,
            "files": records,
            "ok": len(errors) == 0,
            "errors": errors,
        }
        json_path, txt_path = write_reports(report, cfg.get_report_dir())
        report["report_json"] = str(json_path)
        report["report_txt"] = str(txt_path)

        if verbose:
            print(format_summary(report))
            print(f"JSON report     : {json_path}")
            print(f"Text summary    : {txt_path}")

        return report
    finally:
        conn.close()


def run_integrity(
    work_dir: Path | None = None,
    *,
    expected_path: Path | None = None,
    mock_embed: bool = True,
    verbose: bool = True,
) -> int:
    """CLI entry: 0 = all files explained and expected states match."""
    report = run_integrity_benchmark(
        work_dir,
        expected_path=expected_path,
        mock_embed=mock_embed,
        verbose=verbose,
    )
    return 0 if report.get("ok") else 1
